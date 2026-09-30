"""Word (.docx) renderer (python-docx).

Same layout as the PDF: a cover page, a table of contents, numbered h1 sections, tables
with a coloured header row. The table of contents is a Word field that Word fills in when
the document is opened and the field is updated.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from repolens.i18n import t as tr  # `t` is the block type in render()
from repolens.render.inline import segments

BRAND = RGBColor(0x4D, 0x14, 0x8C)
MUTED = RGBColor(0x62, 0x5C, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MONO = "Consolas"
USABLE_WIDTH_CM = 17.0  # A4 width minus the margins


def _shade(cell_or_par, hex_fill: str) -> None:
    """Background colour of a table cell or a paragraph."""
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    props = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    props.append(shd)


def _add_runs(
    par, text: str, size: float | None = None, color: RGBColor | None = None, bold: bool = False
) -> None:
    """Add text with inline `code`, **bold**, and *italic* to a paragraph."""
    for content, style in segments(text):
        run = par.add_run(content)
        if style == "code":
            run.font.name = MONO
            run.font.size = Pt((size or 10.5) - 1)
        elif style == "bold":
            run.bold = True
        elif style == "italic":
            run.italic = True
        if bold:
            run.bold = True
        if size and style != "code":
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def _field(par, instruction: str) -> None:
    """A Word field such as PAGE or TOC; Word computes its value."""
    run = par.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _repeat_header(row) -> None:
    """Repeat this table row at the top of every page the table spans."""
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def render(blocks: list[dict], path: Path) -> Path:
    """Write the blocks to a .docx at `path`: cover, table of contents field, then the body."""
    doc = _new_document()
    state = {"title": "", "h1": 0}  # the title for the page header; h1 numbering
    for block in blocks:
        writer = WRITERS.get(block["t"])
        if writer:
            writer(doc, block, state)
    section = doc.sections[0]
    header = f"{state['title']} · " + tr("Technical Documentation", "Dokumentasi Teknis")
    _add_runs(section.header.paragraphs[0], header, size=8, color=MUTED)
    footer = section.footer.paragraphs[0]
    _add_runs(footer, tr("Page ", "Halaman "), size=8, color=MUTED)
    _field(footer, "PAGE")
    doc.save(path)
    return path


def _new_document():
    """An A4 document with 2 cm margins, Calibri text, and headings in the brand colour."""
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2))
    section.top_margin = section.bottom_margin = Cm(2)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        styles[name].font.color.rgb = BRAND
        styles[name].font.size = Pt(size)
    return doc


# ---- one writer per block type ------------------------------------------------------------


def _title(doc, block: dict, state: dict) -> None:
    """Cover page with metadata, then a table of contents field (Word fills it on update)."""
    state["title"] = block["title"]
    for _ in range(6):
        doc.add_paragraph()
    _add_runs(doc.add_paragraph(), block["subtitle"].upper(), size=11, color=MUTED, bold=True)
    _add_runs(doc.add_paragraph(), block["title"], size=30, color=BRAND, bold=True)
    doc.add_paragraph()
    table = doc.add_table(rows=0, cols=2)
    for key, value in block["meta"]:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(4.5), Cm(12.5)
        _add_runs(cells[0].paragraphs[0], key, size=10, color=MUTED)
        _add_runs(cells[1].paragraphs[0], value, size=10)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(tr("Contents", "Daftar Isi"), level=1)
    _field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    hint = tr(
        "Right-click and choose *Update Field* to refresh the table of contents.",
        "Klik kanan lalu pilih *Update Field* untuk memperbarui daftar isi.",
    )
    _add_runs(doc.add_paragraph(), hint, size=9, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _h1(doc, block: dict, state: dict) -> None:
    state["h1"] += 1
    doc.add_heading(f"{state['h1']}. {block['text']}", level=1)


def _h2(doc, block: dict, state: dict) -> None:
    doc.add_heading(block["text"], level=2)


def _h3(doc, block: dict, state: dict) -> None:
    doc.add_heading(block["text"], level=3)


def _p(doc, block: dict, state: dict) -> None:
    _add_runs(doc.add_paragraph(), block["text"])


def _bullets(doc, block: dict, state: dict) -> None:
    for item in block["items"]:
        _add_runs(doc.add_paragraph(style="List Bullet"), item)


def _note(doc, block: dict, state: dict) -> None:
    warn = block.get("level") == "warn"
    par = doc.add_paragraph()
    _shade(par, "FDF0E7" if warn else "EFE8F8")
    label = tr("Warning: ", "Perhatian: ") if warn else tr("Note: ", "Catatan: ")
    par.add_run(label).bold = True
    _add_runs(par, block["text"])


def _code(doc, block: dict, state: dict) -> None:
    """Monospace lines in one shaded paragraph (line breaks, not separate paragraphs)."""
    par = doc.add_paragraph()
    _shade(par, "F1EFF5")
    par.paragraph_format.space_after = Pt(8)
    lines = block["text"].split("\n")
    for i, line in enumerate(lines):
        run = par.add_run(line)
        run.font.name = MONO
        run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO)
        run.font.size = Pt(8)
        if i < len(lines) - 1:
            run.add_break()


def _table(doc, block: dict, state: dict) -> None:
    """Header row in the brand colour (repeated on every page), zebra rows, relative widths."""
    headers, rows = block["headers"], block["rows"]
    weights = block.get("widths") or [1] * len(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    font_size = 8.5 if len(headers) >= 4 else 9.5
    header_row = table.rows[0]
    _repeat_header(header_row)
    for i, heading in enumerate(headers):
        cell = header_row.cells[i]
        _shade(cell, "4D148C")
        _add_runs(cell.paragraphs[0], heading, size=font_size, color=WHITE, bold=True)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            _add_runs(cells[i].paragraphs[0], str(value), size=font_size)
            if index % 2:
                _shade(cells[i], "F6F5F9")
    for row in table.rows:
        for i, weight in enumerate(weights):
            row.cells[i].width = Cm(USABLE_WIDTH_CM * weight / sum(weights))
    doc.add_paragraph()


WRITERS = {
    "title": _title,
    "h1": _h1,
    "h2": _h2,
    "h3": _h3,
    "p": _p,
    "bullets": _bullets,
    "note": _note,
    "code": _code,
    "table": _table,
}
